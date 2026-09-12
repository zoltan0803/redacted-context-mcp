from __future__ import annotations

import importlib.util
import io
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from redacted_context_mcp import core, documents, server
from redacted_context_mcp.discovery import discover_entities
from redacted_context_mcp.filesystem import RedactedContext
from redacted_context_mcp.limits import OperationBudget, OperationLimitError
from redacted_context_mcp.models import DiscoveryResult, RedactionConfig

HAS_DOCUMENTS = importlib.util.find_spec("markitdown") is not None
TEXT = "quokkaproject database backup procedures for PostgreSQL"


def write_docx(path: Path, text: str = TEXT) -> None:
    from xml.sax.saxutils import escape
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr("word/document.xml", f'<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>{escape(text)}</w:t></w:r></w:p></w:body></w:document>')


def write_pdf(path: Path, text: str = TEXT) -> None:
    content = f"BT /F1 12 Tf 40 700 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for number, value in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{number} 0 obj\n".encode() + value + b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    data += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(data)


class DocumentBoundaryTest(unittest.TestCase):
    def test_plain_text_path_never_imports_optional_library(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            file = root / "notes.txt"
            file.write_text("ordinary text", encoding="utf-8")
            with patch.object(documents.importlib.util, "find_spec", side_effect=AssertionError("optional library loaded")):
                self.assertEqual(RedactedContext(root, RedactionConfig()).read_text(file), "ordinary text")

    def test_missing_extra_is_actionable_and_opt_in(self) -> None:
        with patch.object(documents.importlib.util, "find_spec", return_value=None):
            with self.assertRaisesRegex(SystemExit, r"redacted-context-mcp\[documents\]"):
                RedactedContext(Path.cwd(), RedactionConfig(), documents=True)

    def test_worker_failures_and_timeout_do_not_expose_raw_diagnostics(self) -> None:
        with patch.object(documents, "require_document_support"):
            for code in (1, 2, 3, 4):
                with self.subTest(code=code):
                    result = subprocess.CompletedProcess([], code, stdout=b"private-parser-canary")
                    with patch.object(documents.subprocess, "run", return_value=result):
                        with self.assertRaises(SystemExit) as caught:
                            documents.extract_document(b"dummy", ".pdf")
                    self.assertNotIn("private-parser-canary", str(caught.exception))
            with patch.object(documents.subprocess, "run", side_effect=subprocess.TimeoutExpired("private-parser-canary", 1)):
                with self.assertRaisesRegex(SystemExit, "deadline exceeded"):
                    documents.extract_document(b"dummy", ".pdf")

    def test_input_limit_prevents_worker_launch(self) -> None:
        with patch.object(documents, "require_document_support"), patch.object(documents, "MAX_DOCUMENT_BYTES", 2):
            with patch.object(documents.subprocess, "run", side_effect=AssertionError("worker launched")):
                with self.assertRaisesRegex(SystemExit, "byte limit"):
                    documents.extract_document(b"large", ".pdf")

    def test_zip_expansion_limit_precedes_optional_converter_imports(self) -> None:
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", "x" * 1000)
        with patch.object(documents, "MAX_EXPANDED_BYTES", 999):
            with self.assertRaises(SystemExit) as caught:
                documents.convert_bytes(stream.getvalue(), ".docx")
        self.assertEqual(caught.exception.code, 3)


@unittest.skipUnless(HAS_DOCUMENTS, "optional documents extra is not installed")
class DocumentIntegrationTest(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name).resolve()
        self.config = self.root / ".agent-context-redactor.toml"
        self.config.write_text('[redaction]\nsalt = "document-test-salt"\nterms = ["quokkaproject"]\n', encoding="utf-8")
        self.mcp = server.RedactedContextMcp(root=self.root, config_path=None, mode="balanced", include_private=False, documents=True)

    def test_real_docx_pptx_pdf_and_xlsx_read_redacted(self) -> None:
        from pptx import Presentation
        from openpyxl import Workbook

        write_docx(self.root / "notes.docx")
        write_pdf(self.root / "notes.pdf")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "database backup"
        slide.placeholders[1].text = TEXT
        presentation.save(self.root / "notes.pptx")
        workbook = Workbook()
        workbook.active.append(["description", "action"])
        workbook.active.append([TEXT, "backup"])
        workbook.save(self.root / "notes.xlsx")
        for extension in ("docx", "pptx", "pdf", "xlsx"):
            with self.subTest(extension=extension):
                result = self.mcp.call_tool("redctx_read", {"path": f"notes.{extension}"})
                self.assertFalse(result["isError"], result)
                self.assertNotIn("quokkaproject", json.dumps(result))
                self.assertIn("PostgreSQL", json.dumps(result))
                self.assertIn("[SENSITIVE_", json.dumps(result))

    def test_search_retrieve_bundle_and_cached_resources_share_extraction(self) -> None:
        write_docx(self.root / "notes.docx")
        for name, args in (
            ("redctx_search", {"query": "PostgreSQL"}),
            ("redctx_search", {"query": "PostgreSQL", "regex": True}),
            ("redctx_retrieve", {"query": "database backup"}),
            ("redctx_bundle", {"paths": ["notes.docx"]}),
        ):
            with self.subTest(tool=name, args=args):
                result = self.mcp.call_tool(name, args)
                self.assertFalse(result["isError"], result)
                self.assertIn("PostgreSQL", json.dumps(result))
                self.assertNotIn("quokkaproject", json.dumps(result))
        resource = self.mcp.list_resources({})["resources"][0]
        self.assertEqual(resource["mimeType"], "text/markdown")
        content = self.mcp.read_resource({"uri": resource["uri"]})
        with patch("redacted_context_mcp.filesystem.extract_document", side_effect=AssertionError("cache miss")):
            self.assertEqual(content, self.mcp.read_resource({"uri": resource["uri"]}))
        self.assertNotIn("quokkaproject", repr(self.mcp.cache.entries))

    def test_live_reload_preserves_documents_and_redacts_new_term(self) -> None:
        write_docx(self.root / "notes.docx", "quokkaproject zebra-project backup")
        uri = self.mcp.list_resources({})["resources"][0]["uri"]
        self.assertIn("zebra-project", json.dumps(self.mcp.read_resource({"uri": uri})))
        self.config.write_text('[redaction]\nsalt = "document-test-salt"\nterms = ["quokkaproject", "zebra-project"]\n', encoding="utf-8")
        self.assertNotIn("zebra-project", json.dumps(self.mcp.read_resource({"uri": uri})))
        self.assertTrue(self.mcp.ctx.documents)

    def test_documents_are_excluded_by_default_and_explicit_exclusions_win(self) -> None:
        path = self.root / "notes.docx"
        write_docx(path)
        config = core.load_config(self.root, None)
        self.assertTrue(RedactedContext(self.root, config).is_excluded(path))
        self.assertFalse(self.mcp.ctx.is_excluded(path))
        self.config.write_text(self.config.read_text(encoding="utf-8") + 'exclude_globs = ["*.docx"]\n', encoding="utf-8")
        self.assertEqual(self.mcp.list_resources({})["resources"], [])

    def test_malformed_and_empty_pdf_fail_without_raw_text(self) -> None:
        (self.root / "broken.pdf").write_text("private-parser-canary", encoding="utf-8")
        result = self.mcp.call_tool("redctx_read", {"path": "broken.pdf"})
        self.assertTrue(result["isError"])
        self.assertNotIn("private-parser-canary", json.dumps(result))
        write_pdf(self.root / "empty.pdf", "")
        result = self.mcp.call_tool("redctx_read", {"path": "empty.pdf"})
        self.assertTrue(result["isError"])
        self.assertIn("OCR", json.dumps(result))

    def test_discovery_receives_extracted_text_and_rehydration_reads_documents(self) -> None:
        path = self.root / "notes.docx"
        write_docx(path)
        class Client:
            def extract(inner, *, rel_path, text):
                self.assertIn("quokkaproject", text)
                self.assertIn("PostgreSQL", text)
                return DiscoveryResult(terms=("quokkaproject",))
        result = discover_entities(self.mcp.ctx, paths=[], globs=[], client=Client(), max_files=10, max_chars_per_file=1000, postprocess=False)
        self.assertIn("quokkaproject", result.terms)
        aliases = core.build_rehydration_map(self.mcp.ctx, self.mcp.redactor, budget=OperationBudget(max_files=10))
        self.assertIn("quokkaproject", aliases.values())

    def test_source_budget_is_checked_before_conversion_and_legacy_is_clear(self) -> None:
        path = self.root / "notes.docx"
        write_docx(path)
        with patch("redacted_context_mcp.filesystem.extract_document", side_effect=AssertionError("converter called")):
            with self.assertRaises(OperationLimitError):
                self.mcp.ctx.read_text(path, budget=OperationBudget(max_raw_bytes_per_file=1))
        (self.root / "notes.doc").write_bytes(b"legacy")
        result = self.mcp.call_tool("redctx_read", {"path": "notes.doc"})
        self.assertTrue(result["isError"])
        self.assertIn("Export", json.dumps(result))


if __name__ == "__main__":
    unittest.main()
